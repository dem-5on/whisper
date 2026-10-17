from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from transcriber.session_export import default_export_path, export_session_audio, export_session_docx, export_session_markdown, render_session_markdown
from transcriber.session_store import SessionStore


class SessionExportTests(unittest.TestCase):
    def test_markdown_contains_transcript_translation_and_structured_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("lecture")
            session = store.update(
                session.id,
                state="ready",
                source_transcript="Kia ora, welcome.",
                english_translation="Hello, welcome.",
                outputs={
                    "notes": ["Topic one", "Topic two"], "questions": ["Why?"],
                    "speaker_names": {"speaker_1": "Dr. Mensah"},
                },
            )

            markdown = render_session_markdown(session)
            exported = export_session_markdown(session, Path(directory) / "lecture.md")

            self.assertIn("# Lecture", markdown)
            self.assertIn("Kia ora, welcome.", markdown)
            self.assertIn("## English translation", markdown)
            self.assertIn("- Topic one", markdown)
            self.assertIn("- Why?", markdown)
            self.assertIn("## Speakers", markdown)
            self.assertIn("Speaker 1: Dr. Mensah", markdown)
            self.assertEqual(exported.read_text(encoding="utf-8"), markdown)

    def test_docx_export_holds_transcript_and_refuses_overwrite(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("dictation")
            session = store.update(
                session.id,
                state="ready",
                source_transcript="Export this dictation.",
            )
            target = Path(directory) / "dictation.docx"
            exported = export_session_docx(session, target)

            self.assertEqual(exported, target)
            read_back = Document(str(target))
            self.assertIn(
                "Export this dictation.",
                "\n".join(paragraph.text for paragraph in read_back.paragraphs),
            )
            with self.assertRaises(FileExistsError):
                export_session_docx(session, target)

    def test_default_export_path_avoids_clobbering(self) -> None:
        from unittest.mock import patch

        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("lecture")
            with patch("transcriber.session_export.Path.home", return_value=Path(directory)):
                first = default_export_path(session, ".docx")
                self.assertEqual(first.parent, Path(directory) / "Documents")
                first.parent.mkdir(parents=True, exist_ok=True)
                first.write_bytes(b"existing")
                second = default_export_path(session, ".docx")
            self.assertNotEqual(first, second)
            self.assertTrue(str(second).endswith(".docx"))

    def test_markdown_contains_lecture_analysis_sections(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("lecture")
            session = store.update(
                session.id,
                state="ready",
                source_transcript="Photosynthesis converts light.",
                outputs={
                    "lecture_notes": "# Notes\n- Light matters.",
                    "topic_groups": [{"title": "Intro", "summary": "Start", "start_segment": 0, "end_segment": 0}],
                    "key_terms": [{"term": "Photosynthesis", "definition": "Plants make food."}],
                    "review_questions": ["Why do plants need light?"],
                },
                finished=True,
            )
            markdown = render_session_markdown(session)
            self.assertIn("## Lecture notes", markdown)
            self.assertIn("### Intro", markdown)
            self.assertIn("**Photosynthesis**: Plants make food.", markdown)
            self.assertIn("1. Why do plants need light?", markdown)

    def test_markdown_contains_meeting_record_sections(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("meeting")
            session = store.update(
                session.id,
                state="ready",
                source_transcript="Kwame: Ship Friday.",
                outputs={
                    "speaker_names": {"speaker_1": "Kwame"},
                    "meeting_summary": "Agreed to ship.",
                    "meeting_decisions": ["Ship on Friday"],
                    "action_items": [{"action": "Cut the release", "owner": "Kwame"}],
                },
                finished=True,
            )
            markdown = render_session_markdown(session)
            self.assertIn("- speaker_1: Kwame".replace("speaker_1", "Speaker 1"), markdown)
            self.assertIn("## Meeting summary", markdown)
            self.assertIn("## Decisions", markdown)
            self.assertIn("Ship on Friday", markdown)
            self.assertIn("## Action items", markdown)
            self.assertIn("Cut the release (Kwame)", markdown)

    def test_docx_contains_meeting_record_sections(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("meeting")
            session = store.update(
                session.id,
                state="ready",
                source_transcript="Kwame: Ship Friday.",
                outputs={
                    "meeting_summary": "Agreed to ship.",
                    "meeting_decisions": ["Ship on Friday"],
                    "action_items": [{"action": "Cut the release", "owner": "Kwame"}],
                },
                finished=True,
            )
            target = Path(directory) / "meeting.docx"
            export_session_docx(session, target)
            text = "\n".join(p.text for p in Document(str(target)).paragraphs)
            for expected in ("Meeting summary", "Decisions", "Ship on Friday", "Action items", "Cut the release"):
                self.assertIn(expected, text)

    def test_docx_contains_lecture_analysis_sections(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("lecture")
            session = store.update(
                session.id,
                state="ready",
                source_transcript="Photosynthesis converts light.",
                outputs={
                    "lecture_notes": "Light matters.",
                    "topic_groups": [{"title": "Intro", "summary": "Start"}],
                    "key_terms": [{"term": "Photosynthesis", "definition": "Plants make food."}],
                    "review_questions": ["Why?"],
                },
                finished=True,
            )
            target = Path(directory) / "lecture.docx"
            export_session_docx(session, target)
            text = "\n".join(p.text for p in Document(str(target)).paragraphs)
            for expected in ("Lecture notes", "Intro", "Photosynthesis", "Why?"):
                self.assertIn(expected, text)

    def test_export_never_overwrites_an_existing_file(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("meeting")
            target = Path(directory) / "meeting.md"
            target.write_text("keep this", encoding="utf-8")

            with self.assertRaises(FileExistsError):
                export_session_markdown(session, target)
            self.assertEqual(target.read_text(encoding="utf-8"), "keep this")

    def test_markdown_renders_timestamped_speaker_segments_with_user_name(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(directory)
            meeting = store.create("meeting")
            meeting = store.update(
                meeting.id,
                source_transcript="Hello there.",
                segments=[{
                    "start_seconds": 2.125, "end_seconds": 3.0,
                    "text": "Hello there.", "speaker_id": "speaker_1",
                }],
                outputs={"speaker_names": {"speaker_1": "Ama"}},
            )

            markdown = render_session_markdown(meeting)

            self.assertIn("`00:00:02.125` **Ama:** Hello there.", markdown)

    def test_audio_export_copies_retained_recording_and_refuses_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("meeting")
            recording = b"RIFF" + b"\x00" * 40 + b"audio-data"
            session.audio_path.write_bytes(recording)
            target = Path(directory) / "meeting.wav"

            exported = export_session_audio(session, target)

            self.assertEqual(exported.read_bytes(), recording)
            self.assertEqual(session.audio_path.read_bytes(), recording)
            with self.assertRaises(FileExistsError):
                export_session_audio(session, target)
            self.assertEqual(target.read_bytes(), recording)

    def test_audio_export_rejects_dictation_without_retained_audio(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SessionStore(Path(directory) / "data")
            session = store.create("dictation")

            with self.assertRaises(FileNotFoundError):
                export_session_audio(session, Path(directory) / "dictation.wav")


if __name__ == "__main__":
    unittest.main()
