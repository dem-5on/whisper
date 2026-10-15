"""Human-readable export for locally stored Wicx sessions."""

from __future__ import annotations

import json
from pathlib import Path
import shutil
from .session_store import StoredSession


def render_session_markdown(session: StoredSession) -> str:
    """Render a session without exposing its local audio path or device data."""
    title = {"dictation": "Dictation", "lecture": "Lecture", "meeting": "Meeting"}.get(
        session.mode, "Wicx session"
    )
    transcript = session.source_transcript.strip() or "_(No transcript available.)_"
    if session.segments:
        speaker_names = session.outputs.get("speaker_names", {})
        rendered_segments: list[str] = []
        for segment in session.segments:
            speaker_id = segment.get("speaker_id")
            speaker_name = speaker_names.get(speaker_id) if isinstance(speaker_names, dict) else None
            speaker = speaker_name or (speaker_id.replace("_", " ").title() if speaker_id else "")
            timestamp = _format_timestamp(segment["start_seconds"])
            label = f"**{speaker}:** " if speaker else ""
            rendered_segments.append(f"- `{timestamp}` {label}{segment['text']}")
        transcript = "\n".join(rendered_segments)
    lines = [
        f"# {title}",
        "",
        f"- Recorded: {session.created_at}",
        f"- Status: {session.state}",
        "",
        "## Transcript",
        "",
        transcript,
    ]
    if session.english_translation:
        lines.extend(["", "## English translation", "", session.english_translation.strip()])
    speaker_names = session.outputs.get("speaker_names")
    if isinstance(speaker_names, dict) and speaker_names:
        lines.extend(["", "## Speakers", ""])
        lines.extend(
            f"- {speaker_id.replace('_', ' ').title()}: {name}"
            for speaker_id, name in speaker_names.items()
        )
    generated_outputs = {
        key: value for key, value in session.outputs.items()
        if key not in {"speaker_names", "analysis_state"}
    }
    lecture_notes = generated_outputs.pop("lecture_notes", "")
    topic_groups = generated_outputs.pop("topic_groups", [])
    key_terms = generated_outputs.pop("key_terms", [])
    review_questions = generated_outputs.pop("review_questions", [])
    if isinstance(lecture_notes, str) and lecture_notes.strip():
        lines.extend(["", "## Lecture notes", "", lecture_notes.strip()])
    if isinstance(topic_groups, list) and topic_groups:
        lines.extend(["", "## Topics", ""])
        for topic in topic_groups:
            if not isinstance(topic, dict):
                continue
            title = str(topic.get("title", "")).strip() or "Untitled"
            lines.append(f"### {title}")
            summary = str(topic.get("summary", "")).strip()
            if summary:
                lines.append(summary)
            lines.append("")
    if isinstance(key_terms, list) and key_terms:
        lines.extend(["", "## Key terms", ""])
        for term in key_terms:
            if not isinstance(term, dict):
                continue
            name = str(term.get("term", "")).strip()
            definition = str(term.get("definition", "")).strip()
            if name and definition:
                lines.append(f"- **{name}**: {definition}")
            elif name:
                lines.append(f"- {name}")
        lines.append("")
    if isinstance(review_questions, list) and review_questions:
        lines.extend(["", "## Review questions", ""])
        for index, question in enumerate(review_questions, 1):
            text = question.get("question", "") if isinstance(question, dict) else question
            text = str(text).strip()
            if text:
                lines.append(f"{index}. {text}")
        lines.append("")
    if generated_outputs:
        lines.extend(["", "## Generated outputs", ""])
        for key, value in generated_outputs.items():
            label = str(key).replace("_", " ").strip().capitalize() or "Output"
            lines.extend([f"### {label}", ""])
            if isinstance(value, list):
                lines.extend([f"- {item}" for item in value] or ["_(No items.)_"])
            elif isinstance(value, (dict, bool, int, float)) or value is None:
                lines.extend(["```json", json.dumps(value, ensure_ascii=False, indent=2), "```"])
            else:
                lines.append(str(value))
            lines.append("")
    if session.error:
        lines.extend(["", "## Processing status", "", session.error])
    return "\n".join(lines).rstrip() + "\n"


def _format_timestamp(seconds: float) -> str:
    total_milliseconds = round(seconds * 1000)
    hours, remainder = divmod(total_milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    whole_seconds, milliseconds = divmod(remainder, 1000)
    return f"{hours:02d}:{minutes:02d}:{whole_seconds:02d}.{milliseconds:03d}"


def export_session_markdown(session: StoredSession, output: str | Path) -> Path:
    """Create a Markdown export, refusing to overwrite an existing file."""
    target = Path(output).expanduser()
    content = render_session_markdown(session)
    try:
        with target.open("x", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
    except FileExistsError:
        raise
    except OSError:
        # If writing failed after exclusive creation, remove only the partial
        # file this call created; never touch a pre-existing export.
        if target.exists():
            try:
                target.unlink()
            except OSError:
                pass
        raise
    return target


def default_export_path(session: StoredSession, suffix: str) -> Path:
    """Suggest a non-clobbering export path in the user's Documents folder."""
    documents = Path.home() / "Documents"
    try:
        documents.mkdir(parents=True, exist_ok=True)
    except OSError:
        documents = Path.cwd()
    date = "".join(ch for ch in str(session.created_at)[:10] if ch.isdigit() or ch == "-")
    stem = f"wicx-{session.mode}-{date or 'session'}"
    candidate = documents / f"{stem}{suffix}"
    counter = 2
    while candidate.exists():
        candidate = documents / f"{stem}-{counter}{suffix}"
        counter += 1
    return candidate


def export_session_docx(session: StoredSession, output: str | Path) -> Path:
    """Create a Word (.docx) export, refusing to overwrite an existing file."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError(
            "Word export needs the 'export' extra: pip install '.[export]'"
        ) from exc
    target = Path(output).expanduser()
    if target.exists():
        raise FileExistsError(f"Export already exists: {target}")
    title = {"dictation": "Dictation", "lecture": "Lecture", "meeting": "Meeting"}.get(
        session.mode, "Wicx session"
    )
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(f"Recorded: {session.created_at} · Status: {session.state}")
    document.add_heading("Transcript", level=2)
    if session.segments:
        speaker_names = session.outputs.get("speaker_names", {})
        if not isinstance(speaker_names, dict):
            speaker_names = {}
        for segment in session.segments:
            speaker_id = segment.get("speaker_id")
            speaker_name = speaker_names.get(speaker_id) if speaker_id else None
            speaker = speaker_name or (
                speaker_id.replace("_", " ").title() if speaker_id else ""
            )
            timestamp = _format_timestamp(segment.get("start_seconds", 0))
            paragraph = document.add_paragraph()
            label = f"[{timestamp}] "
            if speaker:
                label += f"{speaker}: "
            run = paragraph.add_run(label)
            run.bold = True
            run.font.size = Pt(9)
            paragraph.add_run(str(segment.get("text", "")))
    else:
        text = session.source_transcript.strip() or "(No transcript available.)"
        document.add_paragraph(text)
    if session.english_translation:
        document.add_heading("English translation", level=2)
        document.add_paragraph(session.english_translation.strip())
    outputs = session.outputs if isinstance(session.outputs, dict) else {}
    notes = outputs.get("lecture_notes", "")
    if isinstance(notes, str) and notes.strip():
        document.add_heading("Lecture notes", level=2)
        for paragraph_text in notes.strip().split("\n"):
            document.add_paragraph(paragraph_text)
    topics = outputs.get("topic_groups", [])
    if isinstance(topics, list) and topics:
        document.add_heading("Topics", level=2)
        for topic in topics:
            if not isinstance(topic, dict):
                continue
            title_text = str(topic.get("title", "")).strip() or "Untitled"
            document.add_heading(title_text, level=3)
            summary = str(topic.get("summary", "")).strip()
            if summary:
                document.add_paragraph(summary)
    terms = outputs.get("key_terms", [])
    if isinstance(terms, list) and terms:
        document.add_heading("Key terms", level=2)
        for term in terms:
            if not isinstance(term, dict):
                continue
            name = str(term.get("term", "")).strip()
            definition = str(term.get("definition", "")).strip()
            if name and definition:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(f"{name}: ")
                run.bold = True
                paragraph.add_run(definition)
            elif name:
                document.add_paragraph(name)
    questions = outputs.get("review_questions", [])
    if isinstance(questions, list) and questions:
        document.add_heading("Review questions", level=2)
        for index, question in enumerate(questions, 1):
            text = question.get("question", "") if isinstance(question, dict) else question
            text = str(text).strip()
            if text:
                document.add_paragraph(f"{index}. {text}")
    if session.error:
        document.add_heading("Processing status", level=2)
        document.add_paragraph(session.error)
    try:
        document.save(str(target))
    except OSError:
        if target.exists():
            try:
                target.unlink()
            except OSError:
                pass
        raise
    return target


def export_session_audio(session: StoredSession, output: str | Path) -> Path:
    """Copy retained session audio, refusing to overwrite an existing file."""
    if session.mode not in {"lecture", "meeting"}:
        raise FileNotFoundError("Audio is not retained for this session")
    source = session.audio_path
    if not source.is_file() or source.stat().st_size <= 44:
        raise FileNotFoundError("The session recording is missing or empty")
    target = Path(output).expanduser()
    try:
        with source.open("rb") as audio, target.open("xb") as exported:
            shutil.copyfileobj(audio, exported, length=1024 * 1024)
    except FileExistsError:
        raise
    except OSError:
        if target.exists():
            try:
                target.unlink()
            except OSError:
                pass
        raise
    return target
